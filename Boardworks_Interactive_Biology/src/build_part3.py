#!/usr/bin/env python3
"""Part 3 – Sensory receptors. Macro-free .pptx."""

from __future__ import annotations

import os
import sys

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(__file__))
from activities import mcq, stepper, true_false_board
from diagrams import OUT as DIAG, generate_all
from pptx_kit import (
    GREEN,
    GREEN_DARK,
    GREEN_SOFT,
    LIGHT_BLUE,
    LIGHT_GREEN,
    LIGHT_RED,
    NAVY,
    NAVY_TEXT,
    ORANGE,
    RED,
    WHITE,
    YELLOW,
    YELLOW_EDGE,
    Deck,
    add_notes,
    add_picture,
    add_rect,
    add_rich_textbox,
    add_textbox,
    set_shape_text,
    title_slide,
)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)


def P(name):
    return os.path.join(DIAG, name)


def build(out_path=None):
    if not os.path.exists(P("skin_block.png")):
        generate_all()

    d = Deck("Sensory receptors")
    title_slide(
        d, "title", "The Nervous System", "Part 3: Sensory receptors",
        "Interactive reconstruction  ·  Flash-free  ·  Original course © Boardworks Ltd 2009",
    )
    add_picture(d.s("title"), P("title_retina.png"), Inches(2.7), Inches(2.05), height=Inches(3.95))

    s = d.add("contents", "Contents")
    d.chrome("contents")
    items = [
        ("the sensory nervous system", "sensory"),
        ("Pacinian corpuscle animation", "pac_0"),
        ("photoreceptors in the retina", "retina"),
        ("rod cells and rhodopsin", "rods"),
        ("cone cells and iodopsin", "cones"),
        ("rods vs cones sorting", "sort"),
        ("knowledge check", "q1"),
    ]
    y = 1.15
    for i, (label, tgt) in enumerate(items, 1):
        d.button(s, Inches(1.6), Inches(y), Inches(6.8), Inches(0.58), f"{i}.  {label}", fill=GREEN if i % 2 else GREEN_DARK, size=16, target=tgt)
        y += 0.68

    s = d.add("sensory", "The sensory nervous system")
    d.chrome("sensory")
    add_rich_textbox(
        s, Inches(0.35), Inches(0.95), Inches(5.4), Inches(5.3),
        [
            ("Receptors", 18, True, GREEN),
            (" are specialized cells that can detect changes in the body’s internal and external environment. Most receptor cells are only sensitive to one type of stimulus.\n\nReceptors convert the energy of the stimulus into the start of a nerve impulse known as a ", 17, False, NAVY_TEXT),
            ("generator potential", 17, True, GREEN),
            (".\n\nFor example, ", 17, False, NAVY_TEXT),
            ("mechanoreceptors", 17, True, GREEN),
            (" detect changes in mechanical energy, such as pressure.", 17, False, NAVY_TEXT),
        ],
    )
    add_picture(s, P("skin_block.png"), Inches(5.7), Inches(1.1), height=Inches(5.2))
    add_textbox(s, Inches(5.7), Inches(6.3), Inches(3.8), Inches(0.4), "mechanoreceptor in the skin", size=12, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    add_notes(s, "Original wording retained. Skin-block diagram rebuilt (Pacinian corpuscle highlighted in later slides).")

    stepper(
        d, "pac", "Mechanoreceptors in the skin",
        [
            {"caption": "Pacinian corpuscle at rest", "image": P("pacinian.png"),
             "body": "Mechanoreceptors within the skin can detect pressure. The Pacinian corpuscle is a pressure receptor in the dermis/hypodermis. It has concentric lamellae of connective tissue around a sensory neurone ending."},
            {"caption": "Pressure deforms the lamellae", "image": P("pacinian_deformed.png"),
             "body": "Pressure deforms the lamellae and stretches the neurone ending in the centre. Stretch-mediated sodium channels in the membrane open."},
            {"caption": "Generator potential", "image": P("pacinian_deformed.png"),
             "body": "Na⁺ enters and the membrane depolarises – a generator potential. The greater the pressure, the greater the deformation and the larger the generator potential (it is graded)."},
            {"caption": "If threshold is reached", "image": P("skin_pacinian.png"),
             "body": "If the generator potential reaches threshold, voltage-gated Na⁺ channels open and an action potential fires along the sensory neurone towards the CNS. This is how mechanical energy is transduced into electrical energy."},
        ],
        nxt="retina",
        notes="Rebuilt Flash play animation of pressure transduction in a Pacinian corpuscle.",
    )

    s = d.add("retina", "Photoreceptors in the retina")
    d.chrome("retina", activity=True)
    add_textbox(s, Inches(0.3), Inches(0.88), Inches(9.4), Inches(0.35),
                "What cells are found in the retina?", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    d.info_box(s, Inches(0.3), Inches(1.3), Inches(3.15), Inches(3.5),
               "The retina is a photosensitive layer at the back of the eye. It contains two types of photoreceptors called rods and cones, which respond to different light intensities.\n\nPress on each label to find out more about the cells of the retina.", size=13)
    add_picture(s, P("retina_cells.png"), Inches(3.4), Inches(1.2), height=Inches(5.15))
    labels = [
        ("cone cell", 7.55, 1.35, "c_cone"),
        ("rod cell", 7.55, 1.95, "c_rod"),
        ("bipolar cell", 7.55, 2.55, "c_bip"),
        ("ganglion cell", 7.55, 3.15, "c_gang"),
        ("optic nerve", 7.55, 3.75, "c_opt"),
    ]
    for text, x, y, tgt in labels:
        d.label_chip(s, Inches(x), Inches(y), Inches(2.05), Inches(0.48), text + " ▶", target=tgt, size=12)
    d.button(s, Inches(7.55), Inches(6.80), Inches(1.5), Inches(0.32), "Reset", fill=ORANGE, size=12, target="retina")
    add_notes(s, "Rebuilt Flash click-label retina activity.")

    retina_info = {
        "c_cone": ("Cone cell", "Photoreceptor for high light intensity and colour. Outer segment is conical. Concentrated in the fovea. Contains iodopsin. Usually a 1:1 link to a bipolar cell – high visual acuity."),
        "c_rod": ("Rod cell", "Photoreceptor for dim light (scotopic vision). Outer segment is rod-shaped with many rhodopsin discs. Many rods synapse with one bipolar cell (retinal convergence) – high sensitivity, lower acuity. Absent from the fovea."),
        "c_bip": ("Bipolar cell", "Connects photoreceptors to ganglion cells. In rod pathways, several rods share a bipolar cell (summation of generator potentials). In cone pathways at the fovea, the link is often one-to-one."),
        "c_gang": ("Ganglion cell", "Its axon forms part of the optic nerve. An action potential is generated in the ganglion cell, not in the rod or cone themselves. The optic disc (blind spot) is where these axons leave the eye."),
        "c_opt": ("Optic nerve", "Bundle of ganglion-cell axons leaving at the blind spot. There are no photoreceptors here, so an image falling on the optic disc is not seen."),
    }
    for key, (title, body) in retina_info.items():
        s = d.add(key, "Photoreceptors in the retina", hidden=True)
        d.chrome(key, show_nav=False)
        add_picture(s, P("retina_cells.png"), Inches(0.3), Inches(1.1), height=Inches(4.7))
        d.info_box(s, Inches(5.3), Inches(1.2), Inches(4.3), Inches(4.5), f"{title}\n\n{body}", size=15)
        d.button(s, Inches(6.3), Inches(5.95), Inches(2.4), Inches(0.4), "Back", fill=GREEN, target="retina")

    s = d.add("rods", "How do rod cells produce impulses?")
    d.chrome("rods")
    add_rich_textbox(
        s, Inches(0.35), Inches(0.95), Inches(6.3), Inches(5.4),
        [
            ("Rod cells allow vision in dim light due to the presence of a pigment called ", 16, False, NAVY_TEXT),
            ("rhodopsin", 16, True, GREEN),
            (", which is found in membrane-bound vesicles.\n\nWhen rhodopsin absorbs light it splits into its constituent parts, ", 16, False, NAVY_TEXT),
            ("opsin", 16, True, GREEN),
            (" and ", 16, False, NAVY_TEXT),
            ("retinal", 16, True, GREEN),
            (". This is called ", 16, False, NAVY_TEXT),
            ("bleaching", 16, True, GREEN),
            (". Low intensity light is sufficient to cause this breakdown.\n\nThe presence of opsin causes a change in the permeability of the rod cell to sodium, which initiates a generator potential. Rhodopsin can reform in the absence of further light stimulation.", 16, False, NAVY_TEXT),
        ],
    )
    add_picture(s, P("rod_cell.png"), Inches(6.7), Inches(0.95), height=Inches(5.4))
    add_textbox(s, Inches(6.55), Inches(1.0), Inches(3.2), Inches(0.7), "vesicles containing rhodopsin", size=12, bold=True, color=GREEN)
    add_notes(s, "Original wording retained. At A-level, bleaching of rhodopsin → opsin + retinal → generator potential is the required sequence. (More advanced courses note that light actually closes Na⁺ channels in rods; the Boardworks 2009 wording is kept here as in the source course.)")

    s = d.add("cones", "How do cone cells produce impulses?")
    d.chrome("cones")
    add_rich_textbox(
        s, Inches(0.35), Inches(0.95), Inches(6.3), Inches(5.4),
        [
            ("Cone cells are sensitive to high light intensities due to the presence of the pigment ", 16, False, NAVY_TEXT),
            ("iodopsin", 16, True, GREEN),
            (".\n\nIn bright light, iodopsin is broken down into its constituent parts, generating an action potential in the ganglion cell.\n\nThere are three different types of cone cell, each containing a different form of iodopsin. Each form of iodopsin absorbs a different wavelength of light – green, blue or red.\n\nThe colour seen depends on the relative degree of stimulation of the three different types of cone cell.", 16, False, NAVY_TEXT),
        ],
    )
    add_picture(s, P("cone_cell.png"), Inches(6.7), Inches(0.95), height=Inches(5.4))
    add_textbox(s, Inches(6.55), Inches(1.0), Inches(3.2), Inches(0.7), "vesicles containing iodopsin", size=12, bold=True, color=GREEN)
    add_notes(s, "Original wording retained. Trichromatic theory: three iodopsin forms, colour from relative stimulation.")

    # Sorting activity – rods vs cones (original drag)
    s = d.add("sort", "Photoreceptors")
    d.chrome("sort", activity=True)
    add_textbox(s, Inches(0.3), Inches(0.88), Inches(9.4), Inches(0.35),
                "Do these statements apply to rods or cones?", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_textbox(s, Inches(0.45), Inches(1.3), Inches(4.3), Inches(0.4), "Click a statement, then place it.", size=14, color=NAVY)
    statements = [
        ("mostly used for night vision", "sort_night"),
        ("contain rhodopsin", "sort_rho"),
        ("contain iodopsin", "sort_iodo"),
        ("concentrated in the fovea", "sort_fovea"),
        ("high sensitivity, lower acuity", "sort_sens"),
        ("three types detect colour", "sort_col"),
        ("many share one bipolar cell", "sort_conv"),
        ("work in bright light", "sort_bright"),
    ]
    y = 1.75
    col = 0
    for text, tgt in statements:
        x = 0.4 if col == 0 else 5.1
        d.button(s, Inches(x), Inches(y), Inches(4.5), Inches(0.5), text, fill=YELLOW, text_color=NAVY, line=GREEN_DARK, size=13, target=tgt)
        col = 1 - col
        if col == 0:
            y += 0.58
    d.button(s, Inches(0.35), Inches(6.80), Inches(1.7), Inches(0.32), "Check", fill=GREEN, size=12, target="sort_ans")
    d.button(s, Inches(2.15), Inches(6.80), Inches(1.4), Inches(0.32), "Reset", fill=ORANGE, size=12, target="sort")
    add_notes(s, "Rebuilt Flash rods/cones sort. Original first card was ‘mostly used for night vision’ (rods).")

    sort_map = {
        "sort_night": ("mostly used for night vision", "rods"),
        "sort_rho": ("contain rhodopsin", "rods"),
        "sort_iodo": ("contain iodopsin", "cones"),
        "sort_fovea": ("concentrated in the fovea", "cones"),
        "sort_sens": ("high sensitivity, lower acuity", "rods"),
        "sort_col": ("three types detect colour", "cones"),
        "sort_conv": ("many share one bipolar cell", "rods"),
        "sort_bright": ("work in bright light", "cones"),
    }
    for key, (text, which) in sort_map.items():
        s = d.add(key, "Photoreceptors", hidden=True)
        d.chrome(key, show_nav=False)
        add_textbox(s, Inches(0.5), Inches(1.2), Inches(9), Inches(0.6), text, size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_textbox(s, Inches(0.5), Inches(1.9), Inches(9), Inches(0.4), "Does this describe rods or cones?", size=16, color=NAVY, align=PP_ALIGN.CENTER)
        d.button(s, Inches(1.5), Inches(2.6), Inches(3.2), Inches(1.4), "rods", fill=GREEN, size=22, target=f"{key}_r")
        d.button(s, Inches(5.3), Inches(2.6), Inches(3.2), Inches(1.4), "cones", fill=RED, size=22, target=f"{key}_c")
        # feedback
        for choice, label in (("r", "rods"), ("c", "cones")):
            fb = f"{key}_{choice}"
            ok = label == which
            s = d.add(fb, "Photoreceptors", hidden=True)
            d.chrome(fb, show_nav=False)
            fill = LIGHT_GREEN if ok else LIGHT_RED
            edge = GREEN_DARK if ok else RED
            box = add_rect(s, Inches(0.8), Inches(1.8), Inches(8.4), Inches(2.6), fill, line=edge, rounded=True)
            msg = f"Correct – {which}.\n\n{text}" if ok else f"Not quite.\n\n‘{text}’ describes {which}, not {label}."
            set_shape_text(box, msg, size=18, color=NAVY, align=PP_ALIGN.CENTER)
            d.button(s, Inches(3.5), Inches(4.9), Inches(3.0), Inches(0.45), "Back to sort", fill=GREEN, target="sort")

    s = d.add("sort_ans", "Photoreceptors – answers", hidden=True)
    d.chrome("sort_ans", show_nav=False)
    rods = [t for t, w in sort_map.values() if w == "rods"]
    cones = [t for t, w in sort_map.values() if w == "cones"]
    left = add_rect(s, Inches(0.4), Inches(1.0), Inches(4.5), Inches(4.8), LIGHT_GREEN, line=GREEN_DARK, rounded=True)
    set_shape_text(left, "RODS\n\n" + "\n".join("• " + t for t in rods), size=16, color=NAVY, align=PP_ALIGN.LEFT)
    right = add_rect(s, Inches(5.1), Inches(1.0), Inches(4.5), Inches(4.8), LIGHT_RED, line=RED, rounded=True)
    set_shape_text(right, "CONES\n\n" + "\n".join("• " + t for t in cones), size=16, color=NAVY, align=PP_ALIGN.LEFT)
    d.button(s, Inches(3.5), Inches(6.1), Inches(3.0), Inches(0.4), "Back", fill=GREEN, target="sort")

    mcq(d, "q1", "Knowledge check",
        "A Pacinian corpuscle produces a generator potential when:",
        ["stretch-mediated sodium channels open and Na⁺ enters",
         "iodopsin is bleached by dim light",
         "voltage-gated potassium channels close permanently",
         "acetylcholine is released from the lamellae"],
        0, "q2",
        "Pressure deforms the ending, stretch-mediated Na⁺ channels open, Na⁺ influx creates a graded generator potential.",
        "Think about transduction of mechanical energy in the Pacinian corpuscle.")
    mcq(d, "q2", "Knowledge check",
        "Which statement about rods is correct?",
        ["They contain iodopsin and give high visual acuity",
         "They contain rhodopsin and are used in dim light",
         "They are concentrated in the fovea",
         "Each rod has a 1:1 connection to a ganglion cell at the fovea"],
        1, "end",
        "Rods use rhodopsin, operate in dim light, show convergence, and are absent from the fovea.",
        "Compare pigment, location and connectivity of rods and cones.")

    s = d.add("end", "End of Part 3")
    d.chrome("end")
    add_textbox(s, Inches(0.5), Inches(2.0), Inches(9), Inches(1.0),
                "Next: Part 4 – The synapse", size=26, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    d.button(s, Inches(3.3), Inches(4.0), Inches(3.4), Inches(0.5), "Back to contents", fill=GREEN, target="contents")

    out_path = out_path or os.path.join(ROOT, "03_Sensory_Receptors.pptx")
    d.save(out_path)
    return out_path, d


if __name__ == "__main__":
    path, deck = build()
    print("wrote", path, "slides", len(deck.prs.slides))
