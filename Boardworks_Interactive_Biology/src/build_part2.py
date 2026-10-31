#!/usr/bin/env python3
"""Part 2 – The nerve impulse. Macro-free .pptx."""

from __future__ import annotations

import os
import sys

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(__file__))
from activities import mcq, stepper, true_false_board
from diagrams import OUT as DIAG, generate_all
from pptx_kit import (
    GREEN,
    GREEN_DARK,
    GREEN_SOFT,
    LIGHT_BLUE,
    LIGHT_GREEN,
    LIGHT_RED,
    NAVY,
    NAVY_TEXT,
    ORANGE,
    WHITE,
    YELLOW,
    YELLOW_EDGE,
    Deck,
    add_notes,
    add_picture,
    add_rect,
    add_rich_textbox,
    add_textbox,
    set_shape_text,
    title_slide,
)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)


def P(name):
    return os.path.join(DIAG, name)


def build(out_path=None, deck=None):
    own = deck is None
    if own and not os.path.exists(P("ap_graph.png")):
        generate_all()

    d = deck or Deck("The nerve impulse")
    title_slide(
        d, "title", "The Nervous System", "Part 2: The nerve impulse",
        "Interactive reconstruction  ·  Flash-free  ·  Original course © Boardworks Ltd 2009",
    )
    add_picture(d.s("title"), P("title_retina.png"), Inches(2.7), Inches(2.05), height=Inches(3.95))

    s = d.add("contents", "Contents")
    d.chrome("contents")
    items = [
        ("what is a nerve impulse?", "impulse"),
        ("resting potential animation", "rp_0"),
        ("resting potential questions", "rpq"),
        ("action potential animation", "apintro"),
        ("ion channels table", "ions"),
        ("action potential graph", "graph"),
        ("all-or-nothing and stimulus strength", "allornothing"),
        ("refractory period", "refract"),
        ("saltatory conduction", "salt_0"),
        ("true or false plenary", "tf"),
    ]
    y = 0.95
    for i, (label, tgt) in enumerate(items, 1):
        d.button(s, Inches(1.5), Inches(y), Inches(7.0), Inches(0.50), f"{i}.  {label}", fill=GREEN if i % 2 else GREEN_DARK, size=15, target=tgt)
        y += 0.55

    # Teaching slide – original wording retained
    s = d.add("impulse", "What is a nerve impulse?")
    d.chrome("impulse")
    add_rich_textbox(
        s, Inches(0.35), Inches(0.95), Inches(5.4), Inches(5.3),
        [
            ("All cells have a difference in electrical charge across the plasma membrane. This is called the ", 16, False, NAVY_TEXT),
            ("potential difference (p.d.)", 16, True, GREEN),
            (". It is measured in millivolts (mV).\n\n", 16, False, NAVY_TEXT),
            ("The normal, resting state of an axon is called the ", 16, False, NAVY_TEXT),
            ("resting potential", 16, True, GREEN),
            (". In this state the p.d. across the axon is ", 16, False, NAVY_TEXT),
            ("–70 mV", 16, True, NAVY),
            (". The membrane is said to be ", 16, False, NAVY_TEXT),
            ("polarized", 16, True, GREEN),
            (".\n\nA nerve impulse, or ", 16, False, NAVY_TEXT),
            ("action potential", 16, True, GREEN),
            (", occurs when the p.d. across an axon is temporarily reversed. The p.d. changes to around ", 16, False, NAVY_TEXT),
            ("+35 mV", 16, True, NAVY),
            (". The membrane is said to be ", 16, False, NAVY_TEXT),
            ("depolarized", 16, True, GREEN),
            (".", 16, False, NAVY_TEXT),
        ],
    )
    add_picture(s, P("axon_electrodes.png"), Inches(5.6), Inches(1.3), width=Inches(4.1))
    add_notes(s, "Original static teaching slide. Wording retained. Diagram rebuilt.")

    stepper(
        d, "rp", "What happens during the resting state?",
        [
            {"caption": "Charge difference", "image": P("resting_charges.png"),
             "body": "An axon is at resting potential when no impulse is present. The inside is negative compared with the outside (about –70 mV). Click Next step to see how this state is maintained."},
            {"caption": "Sodium–potassium pump", "image": P("membrane_rest.png"),
             "body": "The Na⁺/K⁺ pump uses ATP to move 3 Na⁺ out for every 2 K⁺ in. This is active transport. The unequal exchange makes the inside more negative."},
            {"caption": "Membrane permeability", "image": P("membrane_rest.png"),
             "body": "The resting membrane is much more permeable to K⁺ than to Na⁺ because of open potassium leak channels. K⁺ diffuses out, leaving behind negatively charged organic ions (proteins) that cannot leave."},
            {"caption": "Putting it together", "image": P("resting_charges.png"),
             "body": "Pump (3 Na⁺ out, 2 K⁺ in) + K⁺ leak + trapped organic anions = polarised membrane at about –70 mV. Voltage-gated Na⁺ and K⁺ channels are closed."},
        ],
        nxt="rpq",
        notes="Rebuilt Flash play-bar animation as a teacher-paced step sequence with Reset.",
    )

    # Cloze rebuilt from original dropdown activity
    s = d.add("rpq", "Understanding resting potential")
    d.chrome("rpq")
    add_textbox(s, Inches(0.35), Inches(0.9), Inches(9.3), Inches(0.4),
                "What happens during resting potential? Choose the correct ending for each sentence.",
                size=16, color=NAVY)
    prompts = [
        ("The potential difference across the axon is", "rpq1"),
        ("During this state the axon is", "rpq2"),
        ("The inside of the axon is", "rpq3"),
        ("The sodium–potassium pump moves", "rpq4"),
        ("The axon membrane is more permeable to", "rpq5"),
    ]
    y = 1.45
    for text, tgt in prompts:
        d.button(s, Inches(0.55), Inches(y), Inches(8.9), Inches(0.72), text + "  …", fill=YELLOW, text_color=NAVY, line=GREEN_DARK, size=16, target=tgt)
        y += 0.85
    d.button(s, Inches(3.4), Inches(6.0), Inches(3.2), Inches(0.42), "Show all answers", fill=GREEN_DARK, target="rpq_ans")
    add_notes(s, "Rebuilt Flash drop-down cloze. Original S/C (show/check) buttons become Show all answers.")

    cloze = [
        ("rpq1", "The potential difference across the axon is",
         [("-70 mV", True), ("0 mV", False), ("+35 mV", False)],
         "Resting potential is about –70 mV (some texts use –65 to –70 mV)."),
        ("rpq2", "During this state the axon is",
         [("polarised", True), ("depolarised", False), ("at threshold", False)],
         "The membrane has a charge difference, so it is polarised."),
        ("rpq3", "The inside of the axon is",
         [("negative compared with the outside", True), ("positive compared with the outside", False), ("at the same charge as the outside", False)],
         "Organic anions and net ion movement keep the inside negative."),
        ("rpq4", "The sodium–potassium pump moves",
         [("an unequal number of ions (3 Na⁺ out, 2 K⁺ in)", True), ("equal numbers of Na⁺ and K⁺", False), ("only sodium ions", False)],
         "The 3:2 ratio helps generate the resting potential."),
        ("rpq5", "The axon membrane is more permeable to",
         [("potassium ions than sodium ions", True), ("sodium ions than potassium ions", False), ("calcium ions than potassium ions", False)],
         "K⁺ leak channels are open at rest; voltage-gated Na⁺ channels are closed."),
    ]
    for key, stem, choices, why in cloze:
        s = d.add(key, "Understanding resting potential", hidden=True)
        d.chrome(key, show_nav=False)
        add_textbox(s, Inches(0.5), Inches(1.1), Inches(9), Inches(0.7), stem + " …", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        y = 2.0
        for label, ok in choices:
            tgt = f"{key}_ok" if ok else f"{key}_no"
            # unique wrong targets
            if not ok:
                tgt = f"{key}_no_{label[:8].replace(' ', '')}"
            d.button(s, Inches(1.3), Inches(y), Inches(7.4), Inches(0.7), label, fill=LIGHT_BLUE, text_color=NAVY, line=NAVY, size=15, target=tgt)
            y += 0.85
        # create feedback slides inside loop below

    # Feedback for cloze
    for key, stem, choices, why in cloze:
        s = d.add(f"{key}_ok", "Understanding resting potential", hidden=True)
        d.chrome(f"{key}_ok", show_nav=False)
        box = add_rect(s, Inches(0.7), Inches(1.7), Inches(8.6), Inches(2.6), LIGHT_GREEN, line=GREEN_DARK, rounded=True)
        set_shape_text(box, f"Correct.\n\n{why}", size=18, color=NAVY, align=PP_ALIGN.CENTER)
        d.button(s, Inches(3.5), Inches(4.8), Inches(3.0), Inches(0.45), "Back", fill=GREEN, target="rpq")
        for label, ok in choices:
            if ok:
                continue
            tgt = f"{key}_no_{label[:8].replace(' ', '')}"
            s = d.add(tgt, "Understanding resting potential", hidden=True)
            d.chrome(tgt, show_nav=False)
            box = add_rect(s, Inches(0.7), Inches(1.7), Inches(8.6), Inches(2.4), LIGHT_RED, line=ORANGE, rounded=True)
            set_shape_text(box, "Not quite.\n\nLook back at the resting-potential animation and try again.", size=18, color=NAVY, align=PP_ALIGN.CENTER)
            d.button(s, Inches(3.5), Inches(4.7), Inches(3.0), Inches(0.45), "Try again", fill=ORANGE, target=key)

    s = d.add("rpq_ans", "Resting potential – answers", hidden=True)
    d.chrome("rpq_ans", show_nav=False)
    lines = [
        "• p.d. = –70 mV",
        "• the axon is polarised",
        "• inside is negative compared with the outside",
        "• the pump moves 3 Na⁺ out and 2 K⁺ in (unequal numbers)",
        "• the membrane is more permeable to K⁺ than to Na⁺",
    ]
    d.info_box(s, Inches(0.7), Inches(1.3), Inches(8.6), Inches(3.8), "Answer key\n\n" + "\n".join(lines), size=18)
    d.button(s, Inches(3.5), Inches(5.5), Inches(3.0), Inches(0.45), "Back", fill=GREEN, target="rpq")

    # Action potential intro
    s = d.add("apintro", "Action potential")
    d.chrome("apintro")
    d.info_box(
        s, Inches(0.35), Inches(1.0), Inches(3.5), Inches(4.6),
        "A temporary reversal of the potential difference in the axon is caused by a change in the distribution of ions.\n\nPress play on the next screens to see how this change propagates a nerve impulse.",
        size=16,
    )
    add_picture(s, P("motor_neurone.png"), Inches(3.9), Inches(1.5), width=Inches(5.7))
    d.button(s, Inches(3.5), Inches(5.9), Inches(3.0), Inches(0.45), "Play sequence", fill=GREEN, target="apseq_0")

    stepper(
        d, "apseq", "What happens during an action potential?",
        [
            {"caption": "Resting – channels closed", "image": P("membrane_rest.png"),
             "body": "Voltage-gated Na⁺ and K⁺ channels are closed. The membrane is polarised at about –70 mV."},
            {"caption": "Depolarisation – Na⁺ in", "image": P("membrane_depol.png"),
             "body": "A stimulus opens voltage-gated sodium channels. Na⁺ rushes in down its electrochemical gradient. The inside becomes positive (about +35 mV). This is depolarisation."},
            {"caption": "Repolarisation – K⁺ out", "image": P("membrane_repol.png"),
             "body": "Na⁺ channels inactivate. Voltage-gated K⁺ channels open. K⁺ leaves the axon and the inside becomes negative again. This is repolarisation."},
            {"caption": "Return to rest", "image": P("membrane_rest.png"),
             "body": "There is a brief hyperpolarisation (more negative than rest). The Na⁺/K⁺ pump and leak channels restore the resting ion distribution. Local currents then depolarise the next patch of membrane."},
        ],
        nxt="ions",
        notes="Rebuilt Flash play animation of ion movement during an action potential.",
    )

    # Ion channel table
    s = d.add("ions", "What happens during an action potential?")
    d.chrome("ions", activity=True)
    add_textbox(s, Inches(0.3), Inches(0.88), Inches(9.4), Inches(0.4),
                "Which voltage-gated ion channels are open during each stage?", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    headers = ["", "voltage-gated potassium", "voltage-gated sodium"]
    xs = [0.35, 3.15, 6.45]
    ws = [2.7, 3.2, 3.2]
    for x, w, h in zip(xs, ws, headers):
        b = add_rect(s, Inches(x), Inches(1.4), Inches(w), Inches(0.55), GREEN_SOFT, line=GREEN_DARK, rounded=True)
        set_shape_text(b, h, size=13, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    rows = ["resting potential", "depolarisation", "repolarisation"]
    y = 2.05
    for i, row in enumerate(rows):
        b = add_rect(s, Inches(0.35), Inches(y), Inches(2.7), Inches(0.85), GREEN_SOFT, line=GREEN_DARK, rounded=True)
        set_shape_text(b, row, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        d.button(s, Inches(3.15), Inches(y), Inches(3.2), Inches(0.85), "click to reveal", fill=YELLOW, text_color=NAVY, line=GREEN_DARK, size=14, target="ions_ans")
        d.button(s, Inches(6.45), Inches(y), Inches(3.2), Inches(0.85), "click to reveal", fill=YELLOW, text_color=NAVY, line=GREEN_DARK, size=14, target="ions_ans")
        y += 0.98
    d.button(s, Inches(0.35), Inches(6.80), Inches(1.6), Inches(0.32), "Check", fill=GREEN, size=12, target="ions_ans")
    d.button(s, Inches(2.1), Inches(6.80), Inches(1.4), Inches(0.32), "Reset", fill=ORANGE, size=12, target="ions")
    add_notes(s, "Rebuilt Flash open/closed drag table. Click Check for the key.")

    s = d.add("ions_ans", "Ion channels – answers", hidden=True)
    d.chrome("ions_ans", show_nav=False)
    add_textbox(s, Inches(0.3), Inches(0.9), Inches(9.4), Inches(0.4), "Answer key", size=20, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    data = [
        ("", "K⁺ voltage-gated", "Na⁺ voltage-gated"),
        ("resting potential", "closed", "closed"),
        ("depolarisation", "closed", "open"),
        ("repolarisation", "open", "closed"),
    ]
    y = 1.4
    for i, row in enumerate(data):
        fill = GREEN_DARK if i == 0 else (LIGHT_GREEN if i % 2 else LIGHT_BLUE)
        tc = WHITE if i == 0 else NAVY
        for j, cell in enumerate(row):
            b = add_rect(s, Inches(0.4 + j * 3.05), Inches(y), Inches(2.95), Inches(0.7), fill, rounded=True)
            set_shape_text(b, cell, size=14, bold=True, color=tc, align=PP_ALIGN.CENTER)
        y += 0.8
    add_textbox(s, Inches(0.5), Inches(4.8), Inches(9), Inches(1.1),
                "Note: potassium leak channels stay open at rest. This table is about voltage-gated channels only, matching the original Boardworks activity.",
                size=14, color=NAVY)
    d.button(s, Inches(3.5), Inches(6.2), Inches(3.0), Inches(0.4), "Back", fill=GREEN, target="ions")

    # Graph click labels
    s = d.add("graph", "Changes during an action potential")
    d.chrome("graph", activity=True)
    add_textbox(s, Inches(0.3), Inches(0.85), Inches(9.4), Inches(0.35),
                "How does the potential difference change?", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_picture(s, P("ap_graph.png"), Inches(0.2), Inches(1.2), width=Inches(5.6))
    d.info_box(s, Inches(5.9), Inches(1.35), Inches(3.7), Inches(2.2),
               "Press on the numbered graph labels to view a summary of the events during each stage of an action potential.", size=14)
    for i, (lab, tgt) in enumerate([("1 rest", "g1"), ("2 depolarisation", "g2"), ("3 repolarisation", "g3"), ("4 hyperpolarisation", "g4"), ("5 return to rest", "g5")], 1):
        d.button(s, Inches(5.95), Inches(3.65 + (i - 1) * 0.48), Inches(3.55), Inches(0.42), lab, fill=GREEN if i % 2 else GREEN_DARK, size=12, target=tgt)
    add_notes(s, "Rebuilt Flash click-on-graph. Stages 1–5 match the original numbered labels.")

    graph_text = {
        "g1": ("1  Resting potential", "Voltage-gated channels closed. p.d. ≈ –70 mV. Membrane polarised. Na⁺/K⁺ pump and K⁺ leak maintain the gradient."),
        "g2": ("2  Depolarisation", "Threshold reached. Voltage-gated Na⁺ channels open. Na⁺ influx reverses the p.d. towards +35 mV. The membrane is depolarised."),
        "g3": ("3  Repolarisation", "Na⁺ channels inactivate. Voltage-gated K⁺ channels open. K⁺ efflux restores negativity inside the axon."),
        "g4": ("4  Hyperpolarisation", "K⁺ channels are slow to close, so the p.d. briefly becomes more negative than rest (undershoot)."),
        "g5": ("5  Return to resting potential", "Ion channels reset. The pump restores Na⁺ and K⁺ concentrations. The axon is ready for another impulse once the refractory period ends."),
    }
    for key, (title, body) in graph_text.items():
        s = d.add(key, "Changes during an action potential", hidden=True)
        d.chrome(key, show_nav=False)
        num = int(key[1])
        add_picture(s, P(f"ap_graph_{num}.png"), Inches(0.25), Inches(1.0), width=Inches(5.5))
        d.info_box(s, Inches(5.9), Inches(1.2), Inches(3.7), Inches(4.4), f"{title}\n\n{body}", size=15)
        d.button(s, Inches(6.5), Inches(5.85), Inches(2.4), Inches(0.4), "Back to graph", fill=GREEN, target="graph")

    # Axon change quiz (original start-button activity)
    s = d.add("axonq", "How does the axon change?")
    d.chrome("axonq")
    d.info_box(
        s, Inches(1.3), Inches(1.6), Inches(7.4), Inches(3.4),
        "Changes within the axon are essential for the propagation of an action potential.\n\nAnswer these questions about the properties of depolarisation, repolarisation and resting potential.\n\nPress start to begin.",
        size=18,
    )
    d.button(s, Inches(3.7), Inches(5.2), Inches(2.6), Inches(0.5), "start", fill=ORANGE, size=18, target="aq1")
    add_notes(s, "Rebuilt Flash start-quiz. Three MCQs follow.")

    mcq(d, "aq1", "How does the axon change?",
        "During depolarisation, which ions move where?",
        ["Na⁺ ions move into the axon", "K⁺ ions move into the axon", "Na⁺ ions move out of the axon", "Organic anions leave the axon"],
        0, "aq2",
        "Voltage-gated sodium channels open and Na⁺ floods in, making the inside positive.",
        "Think about which channels open first when threshold is reached.")
    mcq(d, "aq2", "How does the axon change?",
        "During repolarisation, the inside of the axon becomes negative again because:",
        ["Na⁺ continues to enter", "K⁺ ions leave the axon", "The membrane becomes impermeable to all ions", "Chloride ions enter in large amounts"],
        1, "aq3",
        "Voltage-gated potassium channels open and K⁺ efflux restores the negative inside.",
        "Repolarisation is caused by potassium leaving, not by sodium still entering.")
    mcq(d, "aq3", "How does the axon change?",
        "Why does a local circuit cause the next region of membrane to depolarise?",
        ["The pump reverses direction", "Positive charge inside spreads to the adjacent region, opening Na⁺ channels there", "Myelin pumps sodium", "The synapse releases calcium along the axon"],
        1, "allornothing",
        "Inward Na⁺ makes that patch positive. Nearby resting membrane is depolarised to threshold and a new action potential is generated – the impulse propagates.",
        "Local currents (local circuits) are the link between one depolarised patch and the next.")

    s = d.add("allornothing", "The all or nothing principle")
    d.chrome("allornothing")
    add_rich_textbox(
        s, Inches(0.35), Inches(0.95), Inches(5.0), Inches(5.2),
        [
            ("For an action potential to be generated, the stimulus must be greater than the ", 16, False, NAVY_TEXT),
            ("threshold value", 16, True, GREEN),
            (".\n\nA stimulus will be below the threshold value if insufficient numbers of sodium channels open, preventing full depolarization of the axon.\n\nOnce the threshold value is reached, the action potential generated is always the same size regardless of the strength of the stimulus. It is an ‘", 16, False, NAVY_TEXT),
            ("all or nothing", 16, True, GREEN),
            ("’ response.", 16, False, NAVY_TEXT),
        ],
    )
    add_picture(s, P("ap_failed.png"), Inches(5.2), Inches(1.15), width=Inches(4.5))
    add_notes(s, "Original static teaching slide. Graph rebuilt.")

    s = d.add("stim", "Different strengths of stimuli")
    d.chrome("stim", activity=True)
    add_textbox(s, Inches(0.3), Inches(0.88), Inches(9.4), Inches(0.35),
                "Changing the strength of the stimulus", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_picture(s, P("stim_medium.png"), Inches(0.25), Inches(1.25), width=Inches(5.7))
    d.info_box(s, Inches(6.1), Inches(1.3), Inches(3.5), Inches(1.6),
               "Adjust the strength of the stimulus to see how this affects the frequency of an action potential.", size=14)
    d.button(s, Inches(6.2), Inches(3.1), Inches(3.3), Inches(0.5), "weak", fill=GREEN_SOFT, text_color=NAVY, line=GREEN_DARK, size=16, target="stim_w")
    d.button(s, Inches(6.2), Inches(3.75), Inches(3.3), Inches(0.5), "medium", fill=GREEN, size=16, target="stim_m")
    d.button(s, Inches(6.2), Inches(4.4), Inches(3.3), Inches(0.5), "strong", fill=GREEN_DARK, size=16, target="stim_s")
    d.info_box(s, Inches(0.35), Inches(5.55), Inches(9.3), Inches(0.95),
               "A stronger stimulus does not make a bigger action potential. It increases the frequency of action potentials.", size=15)
    add_notes(s, "Rebuilt Flash slider. Weak/medium/strong buttons show different AP frequencies.")

    for key, img, cap in (
        ("stim_w", "stim_weak.png", "Weak (just above threshold): low frequency of action potentials."),
        ("stim_m", "stim_medium.png", "Medium: higher frequency. Amplitude of each action potential is unchanged."),
        ("stim_s", "stim_strong.png", "Strong: highest frequency. Still the same amplitude – all or nothing."),
    ):
        s = d.add(key, "Different strengths of stimuli", hidden=True)
        d.chrome(key, show_nav=False)
        add_picture(s, P(img), Inches(0.3), Inches(1.1), width=Inches(6.0))
        d.info_box(s, Inches(6.4), Inches(1.3), Inches(3.2), Inches(3.5), cap, size=16)
        d.button(s, Inches(6.7), Inches(5.2), Inches(2.6), Inches(0.4), "Back", fill=GREEN, target="stim")

    s = d.add("refract", "The axon after an action potential")
    d.chrome("refract", activity=True)
    add_textbox(s, Inches(0.3), Inches(0.88), Inches(9.4), Inches(0.35),
                "What is the refractory period?", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    add_picture(s, P("refractory.png"), Inches(3.3), Inches(1.2), width=Inches(6.3))
    d.info_box(s, Inches(0.3), Inches(1.25), Inches(2.95), Inches(3.5),
               "Once an action potential has passed, the axon becomes inexcitable for a short period of time. Press on each area of the graph to learn more about how the axon excitability changes.", size=13)
    d.button(s, Inches(0.3), Inches(4.9), Inches(2.95), Inches(0.42), "absolute refractory", fill=NAVY, size=12, target="ref_abs")
    d.button(s, Inches(0.3), Inches(5.4), Inches(2.95), Inches(0.42), "relative refractory", fill=GREEN, size=12, target="ref_rel")
    add_notes(s, "Rebuilt Flash click-on-graph for refractory periods.")

    for key, title, body in (
        ("ref_abs", "Absolute refractory period",
         "No new action potential can be generated, however strong the stimulus. Voltage-gated Na⁺ channels are inactivated and must reset. This also forces unidirectional propagation – the region behind the impulse cannot fire again yet."),
        ("ref_rel", "Relative refractory period",
         "A larger-than-normal stimulus can generate an action potential because some Na⁺ channels have recovered but the membrane is still hyperpolarised. Excitability rises back to normal as the resting potential is restored."),
    ):
        s = d.add(key, "The axon after an action potential", hidden=True)
        d.chrome(key, show_nav=False)
        d.info_box(s, Inches(0.7), Inches(1.5), Inches(8.6), Inches(3.4), f"{title}\n\n{body}", size=17)
        d.button(s, Inches(3.5), Inches(5.3), Inches(3.0), Inches(0.45), "Back", fill=GREEN, target="refract")

    stepper(
        d, "salt", "What affects the speed of an impulse?",
        [
            {"caption": "Myelin and nodes", "image": P("saltatory_0.png"),
             "body": "An action potential moves faster along myelinated neurones than those without a myelin sheath.\n\nIon channels at the nodes of Ranvier allow Na⁺ and K⁺ to cross. It is only at these points that an action potential is generated in a myelinated neurone."},
            {"caption": "Jumping from node to node", "image": P("saltatory_1.png"),
             "body": "The action potential appears at a node. Local currents depolarise the next node. This jumping is saltatory conduction and is much quicker than step-by-step conduction in a non-myelinated axon."},
            {"caption": "Next node fires", "image": P("saltatory_2.png"),
             "body": "The impulse ‘jumps’ to the next exposed node. Myelin prevents ion leakage in the internodes, so the local current remains strong enough to bring the distant node to threshold."},
            {"caption": "Axon diameter and temperature", "image": P("saltatory_3.png"),
             "body": "Greater axon diameter → faster impulse (smaller surface area to volume ratio, less ion leak).\n\nHigher temperature → faster impulse (faster ion diffusion), until proteins denature. This is why local cooling slows nerves."},
        ],
        nxt="speed",
        notes="Rebuilt saltatory conduction animation plus the original speed-factor teaching points.",
    )

    s = d.add("speed", "What affects the speed of an impulse?")
    d.chrome("speed")
    add_textbox(s, Inches(0.45), Inches(1.0), Inches(9.1), Inches(0.45),
                "The speed of a nerve impulse is also affected by:", size=18, color=NAVY)
    d.info_box(s, Inches(0.45), Inches(1.55), Inches(9.1), Inches(2.15),
               "Axon diameter – The greater the diameter of the axon the faster the impulse travels.\n\nAxons with a small diameter have a larger surface area to volume ratio than axons with a large diameter. This causes a larger amount of ions to leak out of the axon, making it more difficult for an action potential to propagate.", size=15)
    d.info_box(s, Inches(0.45), Inches(3.9), Inches(9.1), Inches(1.7),
               "Temperature – The higher the temperature, the faster the speed of the impulse. Temperature affects the rate of diffusion of ions across the axon.", size=15)
    add_notes(s, "Original wording retained.")

    true_false_board(
        d, "tf", "True or false?",
        "Are these statements about action potentials true or false?",
        [
            ("Every stimulus detected by the sensory receptors causes an action potential to be generated.", False,
             "The generator potential must reach threshold. Subthreshold stimuli fail to trigger an action potential."),
            ("An action potential is always the same size regardless of the strength of stimuli.", True,
             "All-or-nothing: once threshold is reached, amplitude is constant. Strength is coded by frequency."),
            ("The amplitude of the action potential increases when the stimulus is stronger.", False,
             "Amplitude stays the same. Frequency increases."),
            ("In an unmyelinated neurone an action potential will travel by saltatory conduction.", False,
             "Saltatory conduction requires myelin and nodes of Ranvier. Unmyelinated axons use continuous conduction."),
            ("It is impossible for an action potential to be generated during the absolute refractory period.", True,
             "Na⁺ channels are inactivated. No stimulus, however strong, can open enough of them."),
        ],
        nxt="end",
    )

    s = d.add("end", "End of Part 2")
    d.chrome("end")
    add_textbox(s, Inches(0.5), Inches(2.0), Inches(9), Inches(1.0),
                "Next: Part 3 – Sensory receptors", size=26, bold=True, color=GREEN, align=PP_ALIGN.CENTER)
    d.button(s, Inches(3.3), Inches(4.0), Inches(3.4), Inches(0.5), "Back to contents", fill=GREEN, target="contents")

    if own:
        out_path = out_path or os.path.join(ROOT, "02_The_Nerve_Impulse.pptx")
        d.save(out_path)
        return out_path, d
    return None, d


if __name__ == "__main__":
    path, deck = build()
    print("wrote", path, "slides", len(deck.prs.slides))
