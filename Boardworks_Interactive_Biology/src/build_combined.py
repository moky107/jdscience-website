#!/usr/bin/env python3
"""Combine the four modules into one master course. Does not overwrite the four PPTX files."""

from __future__ import annotations

import os
import sys

from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

sys.path.insert(0, os.path.dirname(__file__))
from diagrams import OUT as DIAG, generate_all
import build_part1
import build_part2
import build_part3
import build_part4
from pptx_kit import (
    GREEN,
    GREEN_DARK,
    GREEN_HEADER,
    GREEN_SOFT,
    GREY_DARK,
    NAVY,
    NAVY_TEXT,
    PURPLE,
    WHITE,
    Deck,
    add_notes,
    add_oval,
    add_picture,
    add_rect,
    add_textbox,
    set_shape_text,
    title_slide,
    FOOTER_NOTE,
    SLIDE_H,
)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT_NAME = "JD_Science_Complete_Interactive_Nervous_System_Course.pptx"


def P(name):
    return os.path.join(DIAG, name)


def add_menu(d: Deck) -> None:
    s = d.add("home", "Course menu", kind="menu")
    add_rect(s, Inches(0), Inches(0), Inches(10), SLIDE_H, WHITE)
    add_rect(s, Inches(0), Inches(0), Inches(10), Inches(1.05), GREEN_HEADER)
    add_textbox(
        s, Inches(0.4), Inches(0.18), Inches(9.2), Inches(0.4),
        "A2 Biology  ·  JD Science", size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
    )
    add_textbox(
        s, Inches(0.4), Inches(0.52), Inches(9.2), Inches(0.45),
        "The Nervous System – complete interactive course",
        size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
    )
    add_rect(s, Inches(0), Inches(0), Inches(0.08), SLIDE_H, PURPLE)
    add_picture(s, P("title_retina.png"), Inches(0.25), Inches(1.25), height=Inches(4.55))
    add_textbox(
        s, Inches(3.6), Inches(1.20), Inches(6.1), Inches(0.7),
        "Choose a module. Home on every later slide returns here. "
        "Start the slide show (F5) before clicking buttons.",
        size=14, color=NAVY,
    )
    modules = [
        ("p1_title", "Part 1  ·  Organisation of the nervous system", GREEN),
        ("p2_title", "Part 2  ·  The nerve impulse", GREEN_DARK),
        ("p3_title", "Part 3  ·  Sensory receptors", GREEN),
        ("p4_title", "Part 4  ·  The synapse", GREEN_DARK),
    ]
    y = 2.05
    for tgt, label, fill in modules:
        d.button(s, Inches(3.6), Inches(y), Inches(5.9), Inches(0.72), label, fill=fill, size=16, target=tgt)
        y += 0.88
    add_textbox(
        s, Inches(3.6), Inches(5.65), Inches(5.9), Inches(0.85),
        "Teaching sequence is unchanged: explanations, diagrams, activities and questions stay in order inside each module. Hidden answer slides are skipped if you only press Space.",
        size=12, color=NAVY_TEXT,
    )
    add_rect(s, Inches(0.08), Inches(7.22), Inches(9.92), Inches(0.04), PURPLE)
    add_textbox(
        s, Inches(0.3), Inches(7.24), Inches(9.4), Inches(0.24),
        FOOTER_NOTE, size=10, color=PURPLE, align=PP_ALIGN.CENTER,
    )
    prev = add_oval(s, Inches(0.16), Inches(6.45), Inches(0.42), Inches(0.42), WHITE, line=PURPLE, line_pt=1.75)
    set_shape_text(prev, "◀", size=14, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
    nxt = add_oval(s, Inches(9.42), Inches(6.45), Inches(0.42), Inches(0.42), PURPLE, line=PURPLE, line_pt=1.0)
    set_shape_text(nxt, "▶", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    d._nav_shapes["home"] = (prev, nxt)
    add_notes(
        s,
        "Course home. Four module buttons jump to each part title. Sequential Next opens Part 1. "
        "Home buttons on later slides return here.",
    )


def add_module_links(d: Deck) -> None:
    """Cross-part Next module buttons on each end slide. Prefix must be empty."""
    d.prefix = ""
    # Stack under each module's existing "Back to contents" button (x=3.3, w=3.4)
    # so the two controls do not overlap.
    extras = [
        ("p1_end", "Next: The nerve impulse", "p2_title", Inches(3.3), Inches(5.65)),
        ("p2_end", "Next: Sensory receptors", "p3_title", Inches(3.3), Inches(4.65)),
        ("p3_end", "Next: The synapse", "p4_title", Inches(3.3), Inches(4.65)),
        ("p4_end", "Course menu", "home", Inches(3.3), Inches(5.25)),
    ]
    for slide_name, label, tgt, x, y in extras:
        d.button(d.s(slide_name), x, y, Inches(3.4), Inches(0.5), label, fill=GREEN_DARK, size=14, target=tgt)


def build(out_path=None):
    if not os.path.exists(P("neurone_structure.png")):
        generate_all()
    d = Deck("Complete nervous system course")
    d.home_target = "home"
    add_menu(d)

    d.prefix = "p1_"
    build_part1.build(deck=d)
    d.prefix = "p2_"
    build_part2.build(deck=d)
    d.prefix = "p3_"
    build_part3.build(deck=d)
    d.prefix = "p4_"
    build_part4.build(deck=d)

    add_module_links(d)

    out_path = out_path or os.path.join(ROOT, OUT_NAME)
    d.save(out_path)
    return out_path, d


if __name__ == "__main__":
    path, deck = build()
    visible = [n for n in deck.order if n not in deck._hidden]
    print("wrote", path)
    print("total slides", len(deck.prs.slides))
    print("visible", len(visible), "hidden", len(deck._hidden))
    print("first visible", visible[:6])
    print("module titles", [n for n in visible if n.endswith("_title") or n == "home"])
