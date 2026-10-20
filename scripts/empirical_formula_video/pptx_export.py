"""Build an editable 16:9 PowerPoint of the completed lesson slides."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from .theme import LOGO_PATH

OUT = Path(__file__).resolve().parents[2] / "content/videos/empirical-formula"

W, H = Inches(13.333), Inches(7.5)
TEAL = RGBColor(0x2E, 0xC4, 0xCE)
NAVY = RGBColor(0x0A, 0x17, 0x31)
CARD = RGBColor(0x12, 0x2A, 0x44)
ROW = RGBColor(0x17, 0x34, 0x56)
INK = RGBColor(0xF4, 0xF7, 0xFB)
MUTED = RGBColor(0x9B, 0xB4, 0xC9)
GOLD = RGBColor(0xE8, 0xC2, 0x47)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0x2A, 0x5C, 0x80)


def rect(slide, l, t, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if min(w, h) < Inches(6) else MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    return s


def add_text(slide, l, t, w, h, text, size=20, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return box


def footer(slide, page, total=18):
    rect(slide, 0, H - Inches(0.42), W, Inches(0.42), RGBColor(0x07, 0x10, 0x22))
    add_text(slide, Inches(0.4), H - Inches(0.36), Inches(8), Inches(0.3),
             "GCSE CHEMISTRY   ·   EMPIRICAL FORMULAE", 11, False, MUTED)
    add_text(slide, Inches(11.2), H - Inches(0.36), Inches(1.8), Inches(0.3),
             f"{page:02d}  /  {total:02d}", 11, True, INK, PP_ALIGN.RIGHT)


def chrome(slide, page):
    rect(slide, 0, 0, W, H, NAVY)
    rect(slide, 0, 0, W, Inches(0.08), TEAL)
    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(0.38), Inches(0.18), Inches(0.32), Inches(0.32))
    add_text(slide, Inches(0.8), Inches(0.2), Inches(3), Inches(0.3), "JD SCIENCE", 14, True, TEAL)
    rect(slide, Inches(0.32), Inches(0.62), Inches(12.7), Inches(6.36), CARD, LINE)
    footer(slide, page)


def heading(slide, title, subtitle):
    add_text(slide, Inches(0.55), Inches(0.78), Inches(12.2), Inches(0.5), title, 28, True, WHITE)
    add_text(slide, Inches(0.55), Inches(1.28), Inches(12.2), Inches(0.36), subtitle, 16, False, MUTED)


def rows(slide, items, gold_idx=()):
    top = Inches(1.78)
    h = Inches(0.95)
    gap = Inches(0.14)
    for i, text in enumerate(items):
        y = top + (h + gap) * i
        fill = RGBColor(0x24, 0x30, 0x16) if i in gold_idx else ROW
        rect(slide, Inches(0.55), y, Inches(12.2), h, fill)
        color = GOLD if i in gold_idx else INK
        add_text(slide, Inches(0.8), y + Inches(0.28), Inches(11.7), Inches(0.45), text, 20, i in gold_idx, color)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    def s(page):
        slide = prs.slides.add_slide(blank)
        chrome(slide, page)
        return slide

    sl = s(1)
    heading(sl, "Empirical Formulae", "Building on relative atomic mass and relative formula mass")
    rows(sl, [
        "What is an empirical formula?",
        "Calculating from masses",
        "Calculating from percentages",
        "Two questions to practise",
    ])

    sl = s(2)
    heading(sl, "What is an empirical formula?", "The simplest whole-number ratio of atoms")
    rows(sl, [
        "Empirical formula  =  simplest ratio",
        "Glucose:  C₆H₁₂O₆",
        "6 : 12 : 6   ÷ 6   =   1 : 2 : 1",
        "Empirical formula:  CH₂O",
    ], gold_idx=(0, 3))

    sl = s(3)
    heading(sl, "Empirical formula: key idea", "Ratio, not necessarily the molecular formula")
    rows(sl, [
        "Molecular formula:  C₆H₁₂O₆",
        "Empirical formula:  CH₂O",
        "Both describe the same elemental proportions",
    ], gold_idx=(1,))

    sl = s(4)
    heading(sl, "Four-step method", "From mass data to a chemical formula")
    rows(sl, [
        "1   Write down the mass of each element",
        "2   Divide each mass by its Ar",
        "3   Divide mole values by the smallest",
        "4   Write the whole-number ratio",
    ])

    sl = s(5)
    heading(sl, "Worked example 1", "Empirical formula from masses")
    rows(sl, [
        "Mg:  2.4 g     Ar = 24     2.4 ÷ 24 = 0.10 mol     0.10 ÷ 0.10 = 1",
        "O:   1.6 g     Ar = 16     1.6 ÷ 16 = 0.10 mol     0.10 ÷ 0.10 = 1",
        "Ratio  Mg : O  =  1 : 1",
        "Empirical formula  =  MgO",
    ], gold_idx=(3,))

    sl = s(8)
    heading(sl, "Practice question 1", "Pause and try it yourself")
    rows(sl, [
        "Sodium:  4.6 g      Oxygen:  1.6 g",
        "Ar(Na) = 23          Ar(O) = 16",
        "What is the empirical formula?",
        "PAUSE THE VIDEO",
    ], gold_idx=(3,))

    sl = s(9)
    heading(sl, "Practice 1  ·  Answer", "Let’s check your working")
    rows(sl, [
        "Na:  4.6 ÷ 23 = 0.20 mol     0.20 ÷ 0.10 = 2",
        "O:   1.6 ÷ 16 = 0.10 mol     0.10 ÷ 0.10 = 1",
        "Ratio  Na : O  =  2 : 1",
        "Empirical formula  =  Na₂O",
    ], gold_idx=(3,))

    sl = s(11)
    heading(sl, "Worked example 2", "Empirical formula from percentages")
    rows(sl, [
        "C = 40.0%     H = 6.7%     O = 53.3%",
        "Assume a 100 g sample",
        "40.0 g C      6.7 g H      53.3 g O",
        "Ar  C = 12,   H = 1,   O = 16",
    ], gold_idx=(1,))

    sl = s(12)
    heading(sl, "Worked example 2  ·  Working", "Divide each mass by Ar, then by the smallest")
    rows(sl, [
        "C:  40.0 ÷ 12 = 3.33 mol     3.33 ÷ 3.33 ≈ 1",
        "H:   6.7 ÷ 1 = 6.70 mol      6.70 ÷ 3.33 ≈ 2",
        "O:  53.3 ÷ 16 = 3.33 mol     3.33 ÷ 3.33 ≈ 1",
        "Empirical formula  =  CH₂O",
    ], gold_idx=(3,))

    sl = s(14)
    heading(sl, "Practice question 2", "Pause and work through the calculation")
    rows(sl, [
        "Carbon = 52.2%     Hydrogen = 13.0%     Oxygen = 34.8%",
        "Ar:  C = 12,   H = 1,   O = 16",
        "Calculate its empirical formula",
        "PAUSE THE VIDEO",
    ], gold_idx=(3,))

    sl = s(15)
    heading(sl, "Practice 2  ·  Answer", "Assume 100 g, then divide by Ar and by the smallest")
    rows(sl, [
        "C:  52.2 ÷ 12 = 4.35     4.35 ÷ 2.175 = 2",
        "H:  13.0 ÷ 1 = 13.0      13.0 ÷ 2.175 ≈ 6",
        "O:  34.8 ÷ 16 = 2.175    2.175 ÷ 2.175 = 1",
        "Empirical formula  =  C₂H₆O",
    ], gold_idx=(3,))

    sl = s(18)
    heading(sl, "Well done!", "You have completed both practice questions")
    rows(sl, [
        "From masses:  Na₂O",
        "From percentages:  C₂H₆O",
        "JDScience  ·  Learn Smarter. Revise Better. Achieve More.",
    ], gold_idx=(0, 1))

    out = OUT / "jdscience-empirical-formulae.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    print("wrote", build())
